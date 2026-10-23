"""Stream large FeelWM result sheets when the in-memory artifact writer is too large."""

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo
from openpyxl.utils import get_column_letter


RUN = sys.argv[1]
assert RUN in ("FeelWM_baseline", "FeelWM_per_leg")
HERE = Path(__file__).resolve().parent
SOURCE = HERE / "prepared_other" / RUN
DEST = HERE.parent / "outputs" / "01a10e72-e696-7bc0-a394-4bc22e3103ac"
DEST.mkdir(parents=True, exist_ok=True)
TARGET = DEST / f"{RUN}_all_test_times_and_legs.xlsx"

SHEETS = (
    ("By test and leg", "by_test", "ByTest"),
    ("Joint by test", "joint_by_test", "JointTest"),
    ("Force each horizon", "force_horizon", "ForcePoint"),
    ("Joint each horizon", "joint_horizon", "JointPoint"),
    ("How to read", "notes", None),
)

wb = Workbook(write_only=True)
header_font = Font(name="Aptos", size=10, color="FFFFFF", bold=True)
data_font = Font(name="Aptos", size=10, color="17212F")
header_fill = PatternFill(fill_type="solid", fgColor="17365D")
header_alignment = Alignment(vertical="center", wrap_text=True)
data_alignment = Alignment(vertical="center")
note_alignment = Alignment(vertical="center", wrap_text=True)

for title, stem, suffix in SHEETS:
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    if stem != "notes":
        ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 21 if stem != "notes" else 23
    ws.column_dimensions["B"].width = 30 if stem != "notes" else 100
    source = SOURCE / (stem + ".jsonl")
    count = 0
    headers = None
    with source.open(encoding="utf-8") as stream:
        for line in stream:
            row = json.loads(line)
            count += 1
            if headers is None:
                headers = row
                for index in range(3, len(headers) + 1):
                    ws.column_dimensions[get_column_letter(index)].width = 18
                cells = []
                for value in row:
                    cell = WriteOnlyCell(ws, value=value)
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = header_alignment
                    cells.append(cell)
                ws.append(cells)
                continue
            cells = []
            for index, value in enumerate(row):
                cell = WriteOnlyCell(ws, value=value)
                cell.font = data_font
                cell.alignment = note_alignment if stem == "notes" else data_alignment
                if stem != "notes" and isinstance(value, (int, float)) and not isinstance(value, bool):
                    cell.number_format = "0" if index == 2 else "0.00"
                cells.append(cell)
            ws.append(cells)
    if suffix:
        table = Table(displayName=RUN.replace("_", "") + suffix,
                      ref=f"A1:{get_column_letter(len(headers))}{count}")
        table.tableColumns = [TableColumn(id=i, name=name) for i, name in enumerate(headers, 1)]
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
    print(title, count - 1, "rows", flush=True)

wb.save(TARGET)
print("SAVED", TARGET, flush=True)
